"""Build the CBS Talks partnership workbook (spreadsheet MVP of docs/lean-spec.md).

Run:  python3 tools/build_workbook.py  -> template/CBS-Talks-Partnerships.xlsx
The workbook is the deliverable; this script only exists so it can be rebuilt
consistently if the specification changes.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment

OUT = "template/CBS-Talks-Partnerships.xlsx"
N = 500          # data rows prepared per table
FONT = "Arial"

INPUT_FILL = PatternFill("solid", fgColor="FFF8DC")   # light yellow: type here
CALC_FILL = PatternFill("solid", fgColor="EEF1F5")    # grey: calculated, don't edit
HEAD_FILL = PatternFill("solid", fgColor="1F3F73")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
CALC_HEAD_FILL = PatternFill("solid", fgColor="5A6170")
BASE = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=16, bold=True, color="1F3F73")
H2 = Font(name=FONT, size=12, bold=True, color="1F3F73")
MUTED = Font(name=FONT, size=9, italic=True, color="5A6170")
RED = Font(name=FONT, size=10, bold=True, color="B42F2F")
THIN = Side(style="thin", color="DDE2E9")

wb = Workbook()
wb.remove(wb.active)


def name(nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


# ---------------------------------------------------------------- Lists
LISTS = {
    "L_Industry": ["Professional services & consulting", "Finance, banking & insurance", "Technology & software",
                   "Energy & utilities", "Life sciences & healthcare", "Consumer goods & retail",
                   "Shipping, logistics & transport", "Manufacturing & industrials",
                   "Media, marketing & communications", "Investors & startup ecosystem",
                   "Public sector, NGOs & associations", "Hospitality, food & venues", "Other"],
    "L_Size": ["<50", "50-249", "250-999", "1000+", "Unknown"],
    "L_Fit": [0, 1, 2],
    "L_YesNo": ["Yes", "No"],
    "L_YesNoUnknown": ["Yes", "No", "Unknown"],
    "L_EvidenceType": ["None found (checked)", "Sponsors student organisations", "Sponsors university activities",
                       "Sponsors similar events", "Previous CBS Talks partner", "Previous CBS Talks spend"],
    "L_Label": ["Verified fact", "Source-backed inference", "Unknown"],
    "L_Criterion": ["F1 Student pull", "F2 Content fit", "F3 Reason to work with us", "F5 Long-term potential",
                    "Access", "General"],
    "L_Relationship": ["Named", "Champion"],
    "L_Type": ["Speaker", "Recruitment", "Event sponsor", "Recurring sponsor", "Annual partner", "Knowledge",
               "Content", "Networking", "In-kind"],
    "L_Stage": ["Planned", "Contacted", "Meeting", "Proposal sent", "Negotiation", "Won", "Lost", "Nurture"],
    "L_Kind": ["Cash", "In-kind", "Non-cash"],
    "L_Satisfaction": ["Satisfied", "Neutral", "Unsatisfied", "Unknown"],
    "L_Lost": ["No budget", "Chose another organisation", "Not relevant for them", "Bad timing", "No response",
               "Terms not agreed", "We withdrew", "Other"],
    "L_Channel": ["Email", "LinkedIn", "Call", "Meeting", "Event", "Other"],
    "L_Direction": ["Sent by us", "Received from them", "Two-way"],
    "L_Outcome": ["Positive", "Neutral", "Negative", "No reply"],
}
ROLES = [("Campus recruiter / early careers", "Yes"), ("University relations", "Yes"), ("Talent acquisition", "Yes"),
         ("Employer branding", "Yes"), ("HR manager / director", "Yes"), ("Marketing", "Yes"),
         ("Partnerships / sponsorship", "Yes"), ("Communications", "Yes"), ("CSR / sustainability", "Yes"),
         ("Events / sales (venues)", "Yes"), ("Community / ecosystem", "Yes"),
         ("Subject-matter expert / speaker", "Yes"), ("Founder / CEO (small company)", "Yes"),
         ("Country manager / senior executive", "Yes"), ("Executive assistant (routing only)", "No"),
         ("Other", "No")]

# ---------------------------------------------------------------- helpers

def make_table(title, cols, note):
    """cols: list of dicts {h, w, calc, dv, fmt, f (formula template with {r})}. Returns ws, letter map."""
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = "1F3F73"
    ws["A1"] = note
    ws["A1"].font = MUTED
    ws.row_dimensions[1].height = 30
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(len(cols), 12))
    L = {}
    for i, c in enumerate(cols, start=1):
        col = get_column_letter(i)
        L[c["h"]] = col
        cell = ws.cell(row=2, column=i, value=c["h"])
        cell.font = HEAD_FONT
        cell.fill = CALC_HEAD_FILL if c.get("calc") else HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[col].width = c.get("w", 16)
        if c.get("tip"):
            cell.comment = Comment(c["tip"], "CBS Talks")
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = "B3"
    return ws, L


def fill_rows(ws, cols, L, first=3, last=N + 2):
    for i, c in enumerate(cols, start=1):
        col = get_column_letter(i)
        for r in range(first, last + 1):
            cell = ws[f"{col}{r}"]
            cell.font = BASE
            cell.fill = CALC_FILL if c.get("calc") else INPUT_FILL
            cell.border = Border(bottom=THIN)
            if c.get("fmt"):
                cell.number_format = c["fmt"]
            if c.get("f"):
                cell.value = c["f"].format(r=r, **L)
        if c.get("dv"):
            dv = DataValidation(type="list", formula1=c["dv"], allow_blank=True)
            dv.error = "Choose a value from the list."
            dv.errorTitle = "Not in the list"
            ws.add_data_validation(dv)
            dv.add(f"{col}{first}:{col}{last}")


def rng(sheet, col):
    return f"'{sheet}'!${col}$3:${col}${N + 2}"


DATE = "yyyy-mm-dd"
DKK = '#,##0" DKK";-#,##0" DKK";"-"'

# ---------------------------------------------------------------- Settings (CBS Talks)
st = wb.create_sheet("CBS Talks")
st.sheet_properties.tabColor = "C98500"
st["A1"] = "CBS Talks settings"
st["A1"].font = TITLE
st["A2"] = ("Fill in the yellow cells. The semester dates and target drive the dashboard. "
            "Rules (rows 21-27) come from the lean specification; change them only after the back-test.")
st["A2"].font = MUTED
st.merge_cells("A2:C2")
SETTINGS = [
    ("About CBS Talks", None, None, None),
    ("Mission (one paragraph)", "", "S_Mission", None),
    ("Audience summary", "", None, None),
    ("Typical attendance per event", "", None, "0"),
    ("Study-line mix (e.g. BSc IB 22%, MSc FIN 15%)", "", None, None),
    ("Social and newsletter reach (with date)", "", None, None),
    ("Editorial independence policy", "", None, None),
    ("CBS rules for sponsorships: what is allowed", "", None, None),
    ("CBS rules checked on / with whom", "", None, None),
    ("Category exclusivities currently granted", "", None, None),
    ("Current semester", None, None, None),
    ("Semester name", "Example semester", None, None),
    ("Semester start", "=TODAY()-60", "S_SemStart", DATE),
    ("Semester end", "=TODAY()+120", "S_SemEnd", DATE),
    ("Semester cash target (DKK)", 120000, "S_Target", DKK),
    ("Date used for all flags", "=TODAY()", "S_Today", DATE),
    ("Rules (from the lean specification)", None, None, None),
    ("High Fit threshold (Fit ≥)", 6, "S_FitHigh", "0"),
    ("High Access threshold (Access level ≥)", 2, "S_AccessHigh", "0"),
    ("Max companies with a path in progress", 5, "S_PathCap", "0"),
    ("Renewal window (days before renewal date)", 90, "S_RenewalDays", "0"),
    ("No-reply flag after (days)", 10, "S_NoReplyDays", "0"),
    ("Warm contact counts for (days)", 365, "S_WarmDays", "0"),
    ("Contact re-check after (days)", 183, "S_RecheckDays", "0"),
]
r = 4
for label, val, nm, fmt in SETTINGS:
    if val is None and nm is None:
        st.cell(row=r, column=1, value=label).font = H2
        r += 1
        continue
    st.cell(row=r, column=1, value=label).font = BASE
    c = st.cell(row=r, column=2, value=val)
    c.font = BASE
    c.fill = INPUT_FILL
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if fmt:
        c.number_format = fmt
    if nm:
        name(nm, f"'CBS Talks'!$B${r}")
    r += 1
st["C16"] = "Example values: replace with the real semester dates."
st["C18"] = "Example placeholder target: replace with the real target."
st["C19"] = "Keep =TODAY(). Type a fixed date only to review the sheet as of that day."
for a in ("C16", "C18", "C19"):
    st[a].font = MUTED
st.column_dimensions["A"].width = 46
st.column_dimensions["B"].width = 60
st.column_dimensions["C"].width = 50

# ---------------------------------------------------------------- Team / Packages / Themes / Events
team_cols = [dict(h="Name", w=22), dict(h="Role", w=26), dict(h="Email", w=28),
             dict(h="Active", w=10, dv="=L_YesNo"), dict(h="Term start", w=12, fmt=DATE),
             dict(h="Term end", w=12, fmt=DATE)]
team, TL = make_table("Team", team_cols, "Board members who can own companies and opportunities. Set Active = No at handover; their items then show 'needs new owner'.")
fill_rows(team, team_cols, TL)
name("R_TeamName", rng("Team", "A"))
name("R_TeamActive", rng("Team", "D"))

pk_cols = [dict(h="Package", w=24), dict(h="Price (DKK)", w=14, fmt=DKK), dict(h="What's included", w=60),
           dict(h="Capacity per semester", w=12, fmt="0"), dict(h="Active", w=10, dv="=L_YesNo")]
pk, PL = make_table("Packages", pk_cols, "CBS Talks packages and prices. Proposals use these, never a guess of a company's budget.")
fill_rows(pk, pk_cols, PL)
name("R_Package", rng("Packages", "A"))

th_cols = [dict(h="Theme", w=36), dict(h="Year", w=10), dict(h="Active", w=10, dv="=L_YesNo")]
th, ThL = make_table("Themes", th_cols, "This year's editorial themes. Used when rating F2 Content fit.")
fill_rows(th, th_cols, ThL)
name("R_Theme", rng("Themes", "A"))

ev_cols = [dict(h="Event", w=34), dict(h="Date", w=12, fmt=DATE), dict(h="Theme", w=26, dv="=R_Theme"),
           dict(h="Attendance", w=12, fmt="0"), dict(h="Study-line mix", w=36),
           dict(h="Partner company", w=26, dv="=R_Company"), dict(h="Partner report sent on", w=14, fmt=DATE)]
ev, EvL = make_table("Events", ev_cols, "Every CBS Talks event with attendance. This is the evidence for F1 Student pull and for every pitch.")
fill_rows(ev, ev_cols, EvL)

# ---------------------------------------------------------------- Companies
C = "Companies"
OP, CO, IN, CE, EVD = "Opportunities", "Contacts", "Interactions", "Commercial Evidence", "Evidence"
comp_cols = [
    dict(h="Company", w=26, tip="Unique name. Used as the key in all other tabs."),
    dict(h="Website", w=24), dict(h="CVR", w=11), dict(h="Industry", w=24, dv="=L_Industry"),
    dict(h="Size band (DK, info only)", w=11, dv="=L_Size", tip="Descriptive only. Not used in any rule."),
    dict(h="Owner", w=14, dv="=R_TeamName"),
    dict(h="F1 Student pull", w=8, dv="=L_Fit", tip="0 checked, no sign students care · 1 some study lines / moderate interest · 2 clear demand (our attendance data, poll, student requests). Source required for 2."),
    dict(h="F1 reason", w=30), dict(h="F1 source", w=22),
    dict(h="F2 Content fit", w=8, dv="=L_Fit", tip="0 no link to a current theme · 1 relevant topic, no named speaker or idea · 2 named person or concrete event idea for a current theme. Source required for 2."),
    dict(h="F2 reason", w=30), dict(h="F2 source", w=22),
    dict(h="F3 Reason to work with us", w=8, dv="=L_Fit", tip="0 checked, nothing we offer helps them · 1 some hiring of business profiles or interest in our topics · 2 clear recurring hiring of CBS-type profiles or an active agenda on our themes. Source required for 2."),
    dict(h="F3 reason", w=30), dict(h="F3 source", w=22),
    dict(h="F5 Long-term potential", w=8, dv="=L_Fit", tip="0 one-off reason only · 1 could plausibly repeat · 2 recurring need across years or past multi-year partnership. Source required for 2."),
    dict(h="F5 reason", w=30), dict(h="F5 source", w=22),
    dict(h="Fit reviewed on", w=12, fmt=DATE),
    dict(h="Intro path", w=30, tip="Who can introduce us, and whether they agreed. Filled = Access 2."),
    dict(h="Specific need (Quick Win)", w=24, tip="e.g. 'venue for 19 Nov'. Quick Wins only appear in the work queue when this is filled."),
    dict(h="Path in progress", w=9, dv="=L_YesNo"),
    dict(h="Path action", w=30), dict(h="Path started on", w=12, fmt=DATE),
    dict(h="Notes (≤500 chars, nothing personal)", w=34),
    dict(h="Do not contact", w=9, dv="=L_YesNo"), dict(h="Do-not-contact reason", w=22),
]
ws_c, CL = make_table(C, comp_cols, "One row per company. Yellow = type here; grey = calculated (don't edit). Hover a header for the rating anchors. A company enters the matrix only when F1, F2, F3, F5 are rated and Commercial Evidence has at least one row.")
# letters of other tabs (fixed by column order below)
OPL = {"Company": "A", "Type": "B", "Stage": "C", "Owner": "D", "Contact": "E", "Renewal": "F", "Kind": "G",
       "Package": "H", "Proposed": "I", "Agreed": "J", "NonCash": "K", "Next": "L", "Due": "M", "Start": "N",
       "End": "O", "RenewalDate": "P", "Satisfaction": "Q", "SatReason": "R", "LostReason": "S",
       "NurtureReason": "T", "NurtureUntil": "U", "Closed": "V", "Why": "W", "Open": "X", "Flag": "Y",
       "QKey": "Z", "Check": "AA", "WonKey": "AB", "LostKey": "AC"}
COL = {"Company": "A", "Name": "B", "Rel": "E", "Left": "K", "DNC": "L", "Suitable": "M", "Verified": "J"}
INL = {"Date": "A", "Company": "B", "Contact": "C", "Direction": "F", "Outcome": "H", "Warm": "I"}
CEL = {"Company": "A", "Type": "B"}

def R(sheet, col):
    return rng(sheet, col)

c_calc = [
    dict(h="F4 Track record", w=8, calc=True, tip="From Commercial Evidence: 2 = previous CBS Talks spend or sponsors student orgs · 1 = university activities, similar events or previous non-cash partner · 0 = only 'none found' · blank = no evidence recorded yet.",
         f=('=IF($A{r}="","",IF(COUNTIF(' + R(CE, "A") + ',$A{r})=0,"",'
            'IF(COUNTIFS(' + R(CE, "A") + ',$A{r},' + R(CE, "B") + ',"Previous CBS Talks spend")'
            '+COUNTIFS(' + R(CE, "A") + ',$A{r},' + R(CE, "B") + ',"Sponsors student organisations")>0,2,'
            'IF(COUNTIFS(' + R(CE, "A") + ',$A{r},' + R(CE, "B") + ',"Sponsors university activities")'
            '+COUNTIFS(' + R(CE, "A") + ',$A{r},' + R(CE, "B") + ',"Sponsors similar events")'
            '+COUNTIFS(' + R(CE, "A") + ',$A{r},' + R(CE, "B") + ',"Previous CBS Talks partner")>0,1,0))))')),
    dict(h="Fit (0-10)", w=8, calc=True,
         f='=IF($A{r}="","",IF(COUNT(G{r},J{r},M{r},P{r},AB{r})<5,"",G{r}+J{r}+M{r}+P{r}+AB{r}))'),
    dict(h="Access (0-4)", w=8, calc=True, tip="4 champion or satisfied past partner · 3 positive reply within a year · 2 intro path · 1 named contact in a suitable role · 0 no path.",
         f=('=IF($A{r}="","",IF(COUNTIFS(' + R(CO, "A") + ',$A{r},' + R(CO, "E") + ',"Champion",' + R(CO, "K") + ',"<>Yes",' + R(CO, "L") + ',"<>Yes")'
            '+COUNTIFS(' + R(OP, "A") + ',$A{r},' + R(OP, "C") + ',"Won",' + R(OP, "Q") + ',"Satisfied")>0,4,'
            'IF(COUNTIFS(' + R(IN, "B") + ',$A{r},' + R(IN, "I") + ',"Yes",' + R(IN, "A") + ',">="&(S_Today-S_WarmDays))>0,3,'
            'IF(T{r}<>"",2,IF(COUNTIFS(' + R(CO, "A") + ',$A{r},' + R(CO, "M") + ',"Yes",' + R(CO, "K") + ',"<>Yes",' + R(CO, "L") + ',"<>Yes")>0,1,0)))))')),
    dict(h="Quadrant", w=14, calc=True,
         f=('=IF($A{r}="","",IF(Z{r}="Yes","Do not contact",IF(AC{r}="","Research",'
            'IF(AND(AC{r}>=S_FitHigh,AD{r}>=S_AccessHigh),"Priority",IF(AC{r}>=S_FitHigh,"Build a Path",'
            'IF(AD{r}>=S_AccessHigh,"Quick Win","Low Priority"))))))')),
    dict(h="Last contact", w=12, calc=True, fmt=DATE,
         f='=IF($A{r}="","",IF(_xlfn.MAXIFS(' + R(IN, "A") + ',' + R(IN, "B") + ',$A{r})=0,"",_xlfn.MAXIFS(' + R(IN, "A") + ',' + R(IN, "B") + ',$A{r})))'),
    dict(h="Open opportunities", w=9, calc=True,
         f='=IF($A{r}="","",COUNTIFS(' + R(OP, "A") + ',$A{r},' + R(OP, "X") + ',"Yes"))'),
    dict(h="Checks", w=30, calc=True,
         f=('=IF($A{r}="","",_xlfn.TEXTJOIN(", ",TRUE,'
            'IF(COUNTIF($A$3:$A$' + str(N + 2) + ',$A{r})>1,"Duplicate name",""),'
            'IF(OR(AND(G{r}=2,I{r}=""),AND(J{r}=2,L{r}=""),AND(M{r}=2,O{r}=""),AND(P{r}=2,R{r}="")),"Source missing for a 2",""),'
            'IF(F{r}="","No owner",IF(COUNTIFS(R_TeamName,F{r},R_TeamActive,"No")>0,"Needs new owner","")),'
            'IF(AND(V{r}="Yes",X{r}<>"",X{r}<S_SemStart),"Path overdue",""),'
            'IF(AND(S{r}<>"",S{r}<S_Today-365),"Stale research",""),'
            'IF(AND(AE{r}="Research",COUNTIF(' + R(CE, "A") + ',$A{r})=0),"Add commercial evidence","")))')),
    dict(h="Queue key", w=10, calc=True, fmt="0.00000", tip="Sort key for the work queue (helper column).",
         f=('=IF(OR($A{r}="",Z{r}="Yes"),"",'
            'IF(AND(AE{r}="Priority",AG{r}=0,COUNTIFS(' + R(OP, "A") + ',$A{r},' + R(OP, "Y") + ',"Renewal due")=0),600000+(10-AC{r})*1000+(4-AD{r})*100+ROW()/100000,'
            'IF(V{r}="Yes",700000+IF(X{r}="",0,X{r}-40000)/10+ROW()/100000,'
            'IF(AND(AE{r}="Quick Win",U{r}<>"",AG{r}=0),800000+(10-AC{r})*1000+ROW()/100000,'
            'IF(AND(AE{r}="Build a Path",COUNTIF($V$3:$V$' + str(N + 2) + ',"Yes")<S_PathCap),900000+(10-AC{r})*1000+(4-AD{r})*100+ROW()/100000,"")))))')),
]
comp_all = comp_cols + c_calc
# rebuild sheet headers to include calc columns
wb.remove(ws_c)
ws_c, CL = make_table(C, comp_all, "One row per company. Yellow = type here; grey = calculated (don't edit). Hover a header for the rating anchors. A company enters the matrix only when F1, F2, F3, F5 are rated and Commercial Evidence has at least one row.")
assert CL["F4 Track record"] == "AB" and CL["Fit (0-10)"] == "AC" and CL["Access (0-4)"] == "AD" and CL["Quadrant"] == "AE"
assert CL["Open opportunities"] == "AG" and CL["Intro path"] == "T" and CL["Path in progress"] == "V" and CL["Do not contact"] == "Z"
fill_rows(ws_c, comp_all, CL)
name("R_Company", rng(C, "A"))

# ---------------------------------------------------------------- Commercial Evidence
ce_cols = [dict(h="Company", w=26, dv="=R_Company"), dict(h="Type", w=30, dv="=L_EvidenceType"),
           dict(h="What was observed (or what was checked)", w=46),
           dict(h="Source URL", w=30, tip="Required except for 'Previous CBS Talks …' (write the opportunity instead)."),
           dict(h="Amount (DKK) — only for Previous CBS Talks spend", w=14, fmt=DKK,
                tip="Only amounts CBS Talks actually received, from our own records. Never an estimate."),
           dict(h="Year", w=8), dict(h="Recorded by", w=14, dv="=R_TeamName"), dict(h="Recorded on", w=12, fmt=DATE),
           dict(h="Check", w=30, calc=True,
                f=('=IF(A{r}="","",_xlfn.TEXTJOIN(", ",TRUE,'
                   'IF(AND(E{r}<>"",B{r}<>"Previous CBS Talks spend"),"Amount only allowed for Previous CBS Talks spend",""),'
                   'IF(AND(B{r}="Previous CBS Talks spend",E{r}=""),"Add the amount received",""),'
                   'IF(AND(D{r}="",LEFT(B{r},8)<>"Previous",B{r}<>"None found (checked)"),"Source URL needed","")))'))]
ws_ce, CEL2 = make_table(CE, ce_cols, "Observed facts about spending on student or similar engagement. Never a budget estimate. 'None found (checked)' is a real result: describe what you checked.")
fill_rows(ws_ce, ce_cols, CEL2)

# ---------------------------------------------------------------- Contacts
co_cols = [dict(h="Company", w=26, dv="=R_Company"), dict(h="Name", w=22, tip="Unique. Used to link interactions."),
           dict(h="Title", w=26), dict(h="Role category", w=28, dv="=L_RoleName"),
           dict(h="Relationship", w=11, dv="=L_Relationship", tip="Champion = actively pushes for us inside the company."),
           dict(h="CBS alumnus", w=9, dv="=L_YesNoUnknown"),
           dict(h="Work email (published or given to us)", w=28), dict(h="LinkedIn URL", w=28),
           dict(h="Source", w=24), dict(h="Verified on", w=12, fmt=DATE),
           dict(h="Left company", w=9, dv="=L_YesNo"), dict(h="Do not contact", w=9, dv="=L_YesNo"),
           dict(h="Suitable role", w=9, calc=True,
                f='=IF(B{r}="","",IFERROR(INDEX(L_RoleOK,MATCH(D{r},L_RoleName,0)),"No"))'),
           dict(h="Check", w=26, calc=True,
                f=('=IF(B{r}="","",_xlfn.TEXTJOIN(", ",TRUE,'
                   'IF(COUNTIF($B$3:$B$' + str(N + 2) + ',B{r})>1,"Duplicate name",""),'
                   'IF(I{r}="","Source missing",""),'
                   'IF(OR(J{r}="",J{r}<S_Today-S_RecheckDays),"Re-check still in role","")))'))]
ws_co, COL2 = make_table(CO, co_cols, "People at companies. Work data only: published work email or LinkedIn, never guessed addresses or personal opinions (GDPR).")
fill_rows(ws_co, co_cols, COL2)
assert COL2["Suitable role"] == "M" and COL2["Left company"] == "K" and COL2["Do not contact"] == "L"
name("R_ContactName", rng(CO, "B"))

# ---------------------------------------------------------------- Opportunities
OPEN_STAGES = '"Planned","Contacted","Meeting","Proposal sent","Negotiation"'
op_cols = [dict(h="Company", w=24, dv="=R_Company"), dict(h="Type", w=16, dv="=L_Type"),
           dict(h="Stage", w=14, dv="=L_Stage"), dict(h="Owner", w=14, dv="=R_TeamName"),
           dict(h="Contact", w=20, dv="=R_ContactName"), dict(h="Renewal", w=9, dv="=L_YesNo"),
           dict(h="Value kind", w=10, dv="=L_Kind"), dict(h="Package", w=18, dv="=R_Package"),
           dict(h="Proposed value (DKK)", w=13, fmt=DKK), dict(h="Agreed value (DKK)", w=13, fmt=DKK),
           dict(h="Non-cash / in-kind description", w=26), dict(h="Next action", w=30),
           dict(h="Next action due", w=12, fmt=DATE), dict(h="Start date", w=12, fmt=DATE),
           dict(h="End date", w=12, fmt=DATE), dict(h="Renewal date", w=12, fmt=DATE),
           dict(h="Satisfaction", w=11, dv="=L_Satisfaction"), dict(h="Satisfaction reason", w=26),
           dict(h="Lost reason", w=20, dv="=L_Lost"), dict(h="Nurture reason", w=24),
           dict(h="Nurture until", w=12, fmt=DATE), dict(h="Closed on", w=12, fmt=DATE),
           dict(h="Why this type (one sentence)", w=34),
           dict(h="Open", w=7, calc=True,
                f='=IF(A{r}="","",IF(OR(C{r}="Planned",C{r}="Contacted",C{r}="Meeting",C{r}="Proposal sent",C{r}="Negotiation"),"Yes","No"))'),
           dict(h="Flag", w=18, calc=True,
                f=('=IF(A{r}="","",'
                   'IF(AND(C{r}="Won",P{r}<>"",P{r}<=S_Today+S_RenewalDays,'
                   'COUNTIFS($A$3:$A$' + str(N + 2) + ',A{r},$F$3:$F$' + str(N + 2) + ',"Yes",$X$3:$X$' + str(N + 2) + ',"Yes")=0),"Renewal due",'
                   'IF(X{r}<>"Yes","",'
                   'IF(OR(L{r}="",M{r}=""),"Missing next action",'
                   'IF(M{r}<S_Today,"Overdue",'
                   'IF(AND(C{r}="Contacted",'
                   '_xlfn.MAXIFS(' + R(IN, "A") + ',' + R(IN, "B") + ',A{r},' + R(IN, "F") + ',"Sent by us")>0,'
                   '_xlfn.MAXIFS(' + R(IN, "A") + ',' + R(IN, "B") + ',A{r},' + R(IN, "F") + ',"Sent by us")<=S_Today-S_NoReplyDays,'
                   'MAX(_xlfn.MAXIFS(' + R(IN, "A") + ',' + R(IN, "B") + ',A{r},' + R(IN, "F") + ',"Received from them"),'
                   '_xlfn.MAXIFS(' + R(IN, "A") + ',' + R(IN, "B") + ',A{r},' + R(IN, "F") + ',"Two-way"))'
                   '<_xlfn.MAXIFS(' + R(IN, "A") + ',' + R(IN, "B") + ',A{r},' + R(IN, "F") + ',"Sent by us")),"No reply",'
                   'IF(M{r}<=S_Today+7,"Due this week",""))))))')),
           dict(h="Queue key", w=10, calc=True, fmt="0.00000",
                f=('=IF(OR(A{r}="",Y{r}=""),"",'
                   'IF(Y{r}="Renewal due",100000,IF(Y{r}="Overdue",200000,IF(Y{r}="No reply",300000,'
                   'IF(Y{r}="Due this week",400000,500000))))'
                   '+IF(Y{r}="Renewal due",P{r}-40000,IF(Y{r}="Missing next action",0,M{r}-40000))+ROW()/100000)')),
           dict(h="Check", w=30, calc=True,
                f=('=IF(A{r}="","",_xlfn.TEXTJOIN(", ",TRUE,'
                   'IF(AND(X{r}="Yes",D{r}=""),"No owner",IF(AND(D{r}<>"",COUNTIFS(R_TeamName,D{r},R_TeamActive,"No")>0),"Needs new owner","")),'
                   'IF(AND(OR(C{r}="Proposal sent",C{r}="Negotiation"),H{r}="",I{r}="",K{r}=""),"Add package or proposed value",""),'
                   'IF(AND(C{r}="Won",OR(P{r}="",V{r}="")),"Won needs renewal date and closed on",""),'
                   'IF(AND(C{r}="Won",G{r}="Cash",J{r}=""),"Add agreed value",""),'
                   'IF(AND(C{r}="Lost",OR(S{r}="",V{r}="")),"Lost needs reason and closed on",""),'
                   'IF(AND(C{r}="Nurture",OR(T{r}="",U{r}="")),"Nurture needs reason and date","")))')),
           dict(h="Won key", w=8, calc=True, fmt="0.00000",
                f='=IF(AND(C{r}="Won",V{r}<>"",V{r}>=S_SemStart,V{r}<=S_SemEnd),V{r}+ROW()/100000,"")'),
           dict(h="Lost key", w=8, calc=True, fmt="0.00000",
                f='=IF(AND(C{r}="Lost",V{r}<>"",V{r}>=S_SemStart,V{r}<=S_SemEnd),V{r}+ROW()/100000,"")'),
           ]
ws_op, OPL2 = make_table(OP, op_cols, "One row per proposed partnership. Follow-up is a flag, not a stage. Required: Planned = type, owner, next action + due · Proposal sent = package or value · Won = agreed value, renewal date, closed on · Lost = reason · Nurture = reason + date.")
for k, v in {"Open": "X", "Flag": "Y", "Queue key": "Z", "Check": "AA", "Won key": "AB", "Lost key": "AC",
             "Renewal date": "P", "Satisfaction": "Q", "Closed on": "V"}.items():
    assert OPL2[k] == v, (k, OPL2[k])
fill_rows(ws_op, op_cols, OPL2)

# ---------------------------------------------------------------- Interactions
in_cols = [dict(h="Date", w=12, fmt=DATE), dict(h="Company", w=24, dv="=R_Company"),
           dict(h="Contact", w=20, dv="=R_ContactName"), dict(h="Team member", w=14, dv="=R_TeamName"),
           dict(h="Channel", w=10, dv="=L_Channel"), dict(h="Direction", w=16, dv="=L_Direction"),
           dict(h="Summary (one line)", w=50), dict(h="Outcome", w=10, dv="=L_Outcome"),
           dict(h="Counts as warm", w=9, calc=True,
                f=('=IF(A{r}="","",IF(AND(H{r}="Positive",'
                   'IFERROR(INDEX(' + R(CO, "K") + ',MATCH(C{r},' + R(CO, "B") + ',0)),"No")<>"Yes"),"Yes","No"))'))]
ws_in, INL2 = make_table(IN, in_cols, "Log every contact in under a minute. Last contact, Access and the no-reply flag come from here. Direction: Sent by us / Received from them / Two-way (calls, meetings).")
fill_rows(ws_in, in_cols, INL2)

# ---------------------------------------------------------------- Evidence
evd_cols = [dict(h="Company", w=24, dv="=R_Company"), dict(h="Criterion", w=20, dv="=L_Criterion"),
            dict(h="Statement", w=50), dict(h="Label", w=20, dv="=L_Label"),
            dict(h="Reasoning (for inferences)", w=36), dict(h="Source URL", w=30),
            dict(h="AI drafted", w=8, dv="=L_YesNo"), dict(h="Checked by", w=14, dv="=R_TeamName"),
            dict(h="Checked on", w=12, fmt=DATE),
            dict(h="Status", w=24, calc=True,
                 f=('=IF(C{r}="","",IF(AND(G{r}="Yes",H{r}=""),"Unchecked AI draft: do not use",'
                    'IF(AND(D{r}<>"Unknown",F{r}=""),"Needs source",'
                    'IF(AND(D{r}="Source-backed inference",E{r}=""),"Needs reasoning",'
                    'IF(D{r}="","Choose a label","OK")))))'))]
ws_evd, EVDL = make_table(EVD, evd_cols, "Supporting facts for Fit reasons and briefings. Every item is a Verified fact, a Source-backed inference or an Unknown. AI-drafted items stay unusable until a person has checked the source.")
fill_rows(ws_evd, evd_cols, EVDL)

# ---------------------------------------------------------------- Lists sheet
ls = wb.create_sheet("Lists")
ls.sheet_properties.tabColor = "8A8F99"
ls["A1"] = "Dropdown values. Edit only if the specification changes."
ls["A1"].font = MUTED
col = 1
for nm, vals in LISTS.items():
    letter = get_column_letter(col)
    ls.cell(row=2, column=col, value=nm[2:]).font = BOLD
    for i, v in enumerate(vals, start=3):
        ls.cell(row=i, column=col, value=v).font = BASE
    name(nm, f"'Lists'!${letter}$3:${letter}${2 + len(vals)}")
    ls.column_dimensions[letter].width = 30
    col += 1
lr, lo = get_column_letter(col), get_column_letter(col + 1)
ls.cell(row=2, column=col, value="Role category").font = BOLD
ls.cell(row=2, column=col + 1, value="Counts for Access 1").font = BOLD
for i, (role, ok) in enumerate(ROLES, start=3):
    ls.cell(row=i, column=col, value=role).font = BASE
    ls.cell(row=i, column=col + 1, value=ok).font = BASE
name("L_RoleName", f"'Lists'!${lr}$3:${lr}${2 + len(ROLES)}")
name("L_RoleOK", f"'Lists'!${lo}$3:${lo}${2 + len(ROLES)}")
ls.column_dimensions[lr].width = 34
ls.column_dimensions[lo].width = 18

# ---------------------------------------------------------------- Work Queue
wq = wb.create_sheet("Work Queue", 0)
wq.sheet_properties.tabColor = "D03B3B"
wq["A1"] = "Work Queue"
wq["A1"].font = TITLE
wq["A2"] = ('="As of "&TEXT(S_Today,"d mmm yyyy")&" · Work from the top: renewals → follow-ups → Priority prospects → building paths. '
            'Calculated, don\'t edit: update the Opportunities and Companies tabs instead."')
wq["A2"].font = MUTED
wq.merge_cells("A2:H2")
ROWS_A, ROWS_B = 30, 30
wq["A4"] = "1–2 · Renewals and follow-ups (from Opportunities)"
wq["A4"].font = H2
hdrA = ["#", "Why it's here", "Company", "Type", "Stage", "Owner", "Next action", "Due", "Renewal date", "key"]
for i, h in enumerate(hdrA, start=1):
    c = wq.cell(row=5, column=i, value=h)
    c.font = HEAD_FONT
    c.fill = HEAD_FILL
OPR = lambda col: rng(OP, col)
for k in range(1, ROWS_A + 1):
    r = 5 + k
    wq[f"J{r}"] = f'=IFERROR(SMALL({OPR("Z")},{k}),"")'
    m = f'MATCH($J{r},{OPR("Z")},0)'
    wq[f"A{r}"] = f'=IF($J{r}="","",{k})'
    wq[f"B{r}"] = f'=IF($J{r}="","",INDEX({OPR("Y")},{m}))'
    wq[f"C{r}"] = f'=IF($J{r}="","",INDEX({OPR("A")},{m}))'
    wq[f"D{r}"] = f'=IF($J{r}="","",INDEX({OPR("B")},{m}))'
    wq[f"E{r}"] = f'=IF($J{r}="","",INDEX({OPR("C")},{m}))'
    wq[f"F{r}"] = f'=IF($J{r}="","",IF(INDEX({OPR("D")},{m})="","(no owner)",INDEX({OPR("D")},{m})))'
    wq[f"G{r}"] = f'=IF($J{r}="","",IF(INDEX({OPR("L")},{m})="","(add a next action)",INDEX({OPR("L")},{m})))'
    wq[f"H{r}"] = f'=IF($J{r}="","",IF(INDEX({OPR("M")},{m})="","",INDEX({OPR("M")},{m})))'
    wq[f"I{r}"] = f'=IF($J{r}="","",IF(INDEX({OPR("P")},{m})="","",INDEX({OPR("P")},{m})))'
    for col in "ABCDEFGHI":
        wq[f"{col}{r}"].font = BASE
        wq[f"{col}{r}"].border = Border(bottom=THIN)
    wq[f"H{r}"].number_format = DATE
    wq[f"I{r}"].number_format = DATE
    wq[f"J{r}"].font = Font(name=FONT, size=8, color="A0A6B0")
    wq[f"J{r}"].number_format = "0.00000"
startB = 5 + ROWS_A + 3
wq[f"A{startB - 1}"] = "3–4 · Priority prospects, building paths and quick wins (from Companies)"
wq[f"A{startB - 1}"].font = H2
hdrB = ["#", "Why it's here", "Company", "Fit", "Access", "Owner", "Next step", "Quadrant", "Last contact", "key"]
for i, h in enumerate(hdrB, start=1):
    c = wq.cell(row=startB, column=i, value=h)
    c.font = HEAD_FONT
    c.fill = HEAD_FILL
CR = lambda col: rng(C, col)
for k in range(1, ROWS_B + 1):
    r = startB + k
    wq[f"J{r}"] = f'=IFERROR(SMALL({CR("AH")},{k}),"")'.replace("AH", CL["Queue key"])
    m = f'MATCH($J{r},{CR(CL["Queue key"])},0)'
    wq[f"A{r}"] = f'=IF($J{r}="","",{k})'
    wq[f"B{r}"] = (f'=IF($J{r}="","",IF($J{r}<700000,"Priority: no deal yet",IF($J{r}<800000,"Building a path",'
                   f'IF($J{r}<900000,"Quick win for a need","Path candidate (space for one more)"))))')
    wq[f"C{r}"] = f'=IF($J{r}="","",INDEX({CR("A")},{m}))'
    wq[f"D{r}"] = f'=IF($J{r}="","",INDEX({CR("AC")},{m}))'
    wq[f"E{r}"] = f'=IF($J{r}="","",INDEX({CR("AD")},{m}))'
    wq[f"F{r}"] = f'=IF($J{r}="","",IF(INDEX({CR("F")},{m})="","(no owner)",INDEX({CR("F")},{m})))'
    wq[f"G{r}"] = (f'=IF($J{r}="","",IF($J{r}<700000,"Plan the first ask (partnership type cheat sheet) and create an opportunity",'
                   f'IF($J{r}<800000,IF(INDEX({CR("W")},{m})="","(add a path action)",INDEX({CR("W")},{m})),'
                   f'IF($J{r}<900000,"Need: "&INDEX({CR("U")},{m}),"Decide whether to start a path"))))')
    wq[f"H{r}"] = f'=IF($J{r}="","",INDEX({CR("AE")},{m}))'
    wq[f"I{r}"] = f'=IF($J{r}="","",IF(INDEX({CR("AF")},{m})="","",INDEX({CR("AF")},{m})))'
    for col in "ABCDEFGHI":
        wq[f"{col}{r}"].font = BASE
        wq[f"{col}{r}"].border = Border(bottom=THIN)
    wq[f"I{r}"].number_format = DATE
    wq[f"J{r}"].font = Font(name=FONT, size=8, color="A0A6B0")
    wq[f"J{r}"].number_format = "0.00000"
for col, w in zip("ABCDEFGHIJ", [5, 26, 26, 14, 14, 14, 52, 14, 13, 8]):
    wq.column_dimensions[col].width = w
wq.freeze_panes = "A4"
red_rule = FormulaRule(formula=['OR($B6="Overdue",$B6="No reply",$B6="Missing next action")'], font=RED)
wq.conditional_formatting.add(f"B6:B{5 + ROWS_A}", red_rule)

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 1)
db.sheet_properties.tabColor = "2A78D6"
db["A1"] = "Dashboard"
db["A1"].font = TITLE
db["A2"] = '="Semester: "&S_SemesterLabel&" ("&TEXT(S_SemStart,"d mmm yyyy")&" – "&TEXT(S_SemEnd,"d mmm yyyy")&") · counts and recorded amounts only"'
db["A2"].font = MUTED
db.merge_cells("A2:F2")
rr = 4
def head(text):
    global rr
    db.cell(row=rr, column=1, value=text).font = H2
    rr += 1
def row(label, formula, fmt=None):
    global rr
    db.cell(row=rr, column=1, value=label).font = BASE
    c = db.cell(row=rr, column=2, value=formula)
    c.font = BOLD
    if fmt:
        c.number_format = fmt
    rr += 1

head("Matrix (companies)")
for q in ["Priority", "Build a Path", "Quick Win", "Low Priority", "Research"]:
    row(q, f'=COUNTIF({CR("AE")},"{q}")')
row("Paths in progress (max)", f'=COUNTIF({CR("V")},"Yes")&" / "&S_PathCap')
rr += 1
head("Semester money (recorded amounts only, no estimates)")
row("Cash won this semester", f'=SUMIFS({OPR("J")},{OPR("C")},"Won",{OPR("G")},"Cash",{OPR("V")},">="&S_SemStart,{OPR("V")},"<="&S_SemEnd)', DKK)
won_row = rr - 1
row("Semester target", "=S_Target", DKK)
row("Progress", f'=IF(S_Target=0,"",REPT("█",ROUND(MIN(B{won_row}/S_Target,1)*20,0))&REPT("░",20-ROUND(MIN(B{won_row}/S_Target,1)*20,0))&"  "&TEXT(B{won_row}/S_Target,"0%"))')
row("Cash proposed, not committed (Proposal sent + Negotiation)",
    f'=SUMIFS({OPR("I")},{OPR("C")},"Proposal sent",{OPR("G")},"Cash")+SUMIFS({OPR("I")},{OPR("C")},"Negotiation",{OPR("G")},"Cash")', DKK)
rr += 1
head("Open opportunities by stage")
for s in ["Planned", "Contacted", "Meeting", "Proposal sent", "Negotiation", "Nurture"]:
    row(s, f'=COUNTIF({OPR("C")},"{s}")')
rr += 1
head("Non-cash and in-kind (counted, not summed in DKK)")
row("Non-cash won this semester", f'=COUNTIFS({OPR("C")},"Won",{OPR("G")},"Non-cash",{OPR("V")},">="&S_SemStart,{OPR("V")},"<="&S_SemEnd)')
row("In-kind won this semester", f'=COUNTIFS({OPR("C")},"Won",{OPR("G")},"In-kind",{OPR("V")},">="&S_SemStart,{OPR("V")},"<="&S_SemEnd)')
row("Non-cash or in-kind open", f'=COUNTIFS({OPR("X")},"Yes",{OPR("G")},"Non-cash")+COUNTIFS({OPR("X")},"Yes",{OPR("G")},"In-kind")')
rr += 1
head("Data health")
row("Open opportunities without a next action", f'=COUNTIF({OPR("Y")},"Missing next action")')
row("Companies with a check to fix", f'=COUNTIF({CR("AH".replace("AH", "AH"))},"?*")'.replace(CR("AH"), CR(CL["Checks"])))
row("Opportunities with a check to fix", f'=COUNTIF({OPR("AA")},"?*")')
row("Contacts to re-check", f'=COUNTIF({rng(CO, "N")},"*Re-check*")')
row("Unchecked AI evidence", f'=COUNTIF({rng(EVD, "J")},"Unchecked*")')
row("Paths over the cap", f'=MAX(0,COUNTIF({CR("V")},"Yes")-S_PathCap)')

# right-hand side: won / lost lists and lost reasons
def list_block(top, title, keycol, cols):
    db.cell(row=top, column=4, value=title).font = H2
    for i, h in enumerate([c[0] for c in cols], start=4):
        c = db.cell(row=top + 1, column=i, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
    for k in range(1, 11):
        r = top + 1 + k
        key = f'IFERROR(SMALL({OPR(keycol)},{k}),"")'
        m = f'MATCH({key},{OPR(keycol)},0)'
        for i, (_, src, fmt) in enumerate(cols, start=4):
            c = db.cell(row=r, column=i, value=f'=IF(COUNT({OPR(keycol)})<{k},"",IFERROR(IF(INDEX({OPR(src)},{m})="","",INDEX({OPR(src)},{m})),""))')
            c.font = BASE
            c.border = Border(bottom=THIN)
            if fmt:
                c.number_format = fmt
list_block(4, "Won this semester", "AB", [("Company", "A", None), ("Type", "B", None), ("Agreed (DKK)", "J", DKK),
                                          ("Non-cash / in-kind", "K", None), ("Renewal date", "P", DATE),
                                          ("Satisfaction", "Q", None)])
list_block(17, "Lost this semester", "AC", [("Company", "A", None), ("Type", "B", None), ("Lost reason", "S", None),
                                            ("Closed on", "V", DATE)])
db.cell(row=30, column=4, value="Lost reasons this semester").font = H2
for i, reason in enumerate(LISTS["L_Lost"], start=31):
    db.cell(row=i, column=4, value=reason).font = BASE
    db.cell(row=i, column=5, value=f'=COUNTIFS({OPR("C")},"Lost",{OPR("S")},D{i},{OPR("V")},">="&S_SemStart,{OPR("V")},"<="&S_SemEnd)').font = BOLD
for col, w in zip("ABCDEFGHI", [46, 22, 3, 24, 18, 16, 22, 14, 12]):
    db.column_dimensions[col].width = w
name("S_SemesterLabel", "'CBS Talks'!$B$15")

# ---------------------------------------------------------------- Start here
sh = wb.create_sheet("Start here", 0)
sh.sheet_properties.tabColor = "0CA30C"
lines = [
    ("CBS Talks Partnerships", TITLE),
    ("Spreadsheet MVP of docs/lean-spec.md. Two scores only: Fit (0–10) and Access (0–4). No budget estimates.", MUTED),
    ("", None),
    ("How to use it", H2),
    ("1. Delete the example rows: every row whose company or name ends in '(example)', in all tabs.", BASE),
    ("2. Fill in the 'CBS Talks' tab: semester dates and target, then Team, Packages, Themes and Events.", BASE),
    ("3. Import past and current partners: Companies + Opportunities (Won, with renewal date and satisfaction) + Commercial Evidence.", BASE),
    ("4. Rate companies in 'Companies' (F1, F2, F3, F5 with a reason; a source for every 2) and add at least one Commercial Evidence row.", BASE),
    ("5. Every week (30 min): open 'Work Queue', work it top to bottom, log contacts in 'Interactions', update stages and next actions.", BASE),
    ("", None),
    ("Colours", H2),
    ("Yellow cells: type here.   Grey cells: calculated, don't edit.   Dark-blue headers: input columns.   Grey headers: calculated columns.", BASE),
    ("Hover a header in 'Companies' to see the rating anchors.", BASE),
    ("", None),
    ("The rules in one place", H2),
    ("Fit = F1 Student pull + F2 Content fit + F3 Reason to work with us + F4 Track record + F5 Long-term (each 0–2). F4 comes from Commercial Evidence.", BASE),
    ("Access: 4 champion / satisfied past partner · 3 positive reply within a year · 2 intro path · 1 named contact in a suitable role · 0 no path.", BASE),
    ("Matrix: Priority (Fit ≥ 6, Access ≥ 2) · Build a Path (Fit ≥ 6, Access < 2; max 5 at a time) · Quick Win (Fit < 6, Access ≥ 2; only for a specific need) · Low Priority.", BASE),
    ("Work queue order: renewals → overdue → no reply → due this week → missing next action → Priority without a deal → paths in progress → quick wins → path candidates.", BASE),
    ("", None),
    ("Labels for researched or AI-drafted information (Evidence tab)", H2),
    ("Verified fact = a person opened the source and confirmed it.  Source-backed inference = a conclusion from a cited source, with the reasoning written down.  Unknown = an open question.", BASE),
    ("AI-drafted items stay 'Unchecked AI draft: do not use' until a team member has checked the source.", BASE),
    ("", None),
    ("Cheat sheets (in the repository, docs/cheat-sheets/)", H2),
    ("partnership-type.md: what to propose  ·  who-to-contact.md: which role  ·  conversations-and-outreach.md: questions, objections, outreach and GDPR rules", BASE),
    ("", None),
    ("All example companies, people and amounts are fictional.", RED),
]
for i, (t, f) in enumerate(lines, start=1):
    c = sh.cell(row=i, column=1, value=t)
    if f:
        c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")
sh.column_dimensions["A"].width = 140

# ---------------------------------------------------------------- Example data (fictional; dates relative to today)
def T(days):
    return f"=TODAY(){'+' if days >= 0 else '-'}{abs(days)}"

def put(ws, row, values):
    for i, v in enumerate(values, start=1):
        if v is not None and v != "":
            ws.cell(row=row, column=i, value=v)

team_rows = [["A.K. (example)", "Partnerships Manager", "", "Yes", T(-200), T(160)],
             ["J.H. (example)", "Partnerships Associate", "", "Yes", T(-200), T(160)],
             ["M.S. (example)", "Partnerships Associate", "", "Yes", T(-200), T(160)]]
for i, v in enumerate(team_rows, start=3):
    put(team, i, v)
put(pk, 3, ["Entry package (example)", 8000, "Logo on one event, a short company intro on stage", 6, "Yes"])
put(pk, 4, ["Standard package (example)", 20000, "One sponsored event: intro talk, logo, partner report", 4, "Yes"])
put(pk, 5, ["Strategic package (example)", 45000, "Event series over a semester, case workshop, partner report", 2, "Yes"])
put(th, 3, ["Work after AI (example)", "2026/27", "Yes"])
put(th, 4, ["Fragile Systems (example)", "2026/27", "Yes"])
put(ev, 3, ["The future of graduate work (example)", T(-40), "Work after AI (example)", 180, "BSc IB 25%, MSc FIN 15%, other 60%", "ProServ (example)", T(-35)])

companies = [
    ["ProServ (example)", "proserv.example", "", "Professional services & consulting", "1000+", "J.H. (example)",
     2, "180 attended our event with them", "Events tab", 2, "Partner on stage for Work after AI series", "Talk recording",
     2, "Graduate programme, recurring hiring", "careers page", 2, "Partnered two years running", "Opportunities tab",
     T(-20), "", "", "No", "", "", "", "No", ""],
    ["MegaBrand A/S (example)", "megabrand.example", "", "Consumer goods & retail", "1000+", "M.S. (example)",
     2, "Top result in our audience poll", "Poll results", 1, "Relevant topic, no named speaker", "",
     1, "Hires business graduates, already well known", "", 1, "Could repeat", "",
     T(-30), "", "", "Yes", "Ask board alumni network for an introduction", T(-25), "", "No", ""],
    ["Nordlys Logistics (example)", "nordlys.example", "", "Shipping, logistics & transport", "250-999", "A.K. (example)",
     1, "Relevant to supply-chain students", "", 1, "Supply-chain topic, no named speaker yet", "",
     2, "About 20 graduate postings this year", "jobs page", 2, "Annual graduate programme", "careers page",
     T(-15), "CBS alumnus in their talent team, via A.K. (agreed)", "", "No", "", "", "", "No", ""],
    ["Fjord Ventures (example)", "fjord.example", "", "Investors & startup ecosystem", "<50", "A.K. (example)",
     1, "Some interest among entrepreneurship students", "", 2, "Founder has recorded keynotes on fintech", "talk recording",
     1, "Wants visibility for the founder's agenda", "", 1, "Could repeat", "",
     T(-10), "", "Speaker for the spring keynote", "No", "", "", "", "No", ""],
    ["Kobber Hotels (example)", "kobber.example", "", "Hospitality, food & venues", "50-249", "L.M. (example)",
     0, "No sign students care about the brand", "", 0, "No link to our themes", "",
     1, "Interested in local visibility", "", 1, "Could host again", "",
     T(-12), "", "Venue for an autumn event", "No", "", "", "", "No", ""],
    ["Gamma Software (example)", "gamma.example", "", "Technology & software", "50-249", "M.S. (example)",
     0, "Checked: no student interest found", "", 0, "No link to our themes", "",
     0, "No business-graduate hiring found", "", 0, "No recurring reason", "",
     T(-50), "", "", "No", "", "", "", "No", ""],
]
companies += [
    ["Søby Foods (example)", "soby.example", "", "Consumer goods & retail", "250-999", "J.H. (example)"] + [""] * 21,
    ["Strandvej Studio (example)", "strandvej.example", "", "Media, marketing & communications", "<50", "M.S. (example)"] + [""] * 21,
]
for i, v in enumerate(companies, start=3):
    put(ws_c, i, v)
put(team, 6, ["L.M. (example)", "Board member (left)", "", "No", T(-560), T(-30)])

ce_rows = [
    ["ProServ (example)", "Previous CBS Talks spend", "Paid for the event series last year", "", 40000, "2025", "J.H. (example)", T(-20)],
    ["MegaBrand A/S (example)", "Sponsors student organisations", "Main partner of two CBS student organisations", "https://example.org/partners", "", "2026", "M.S. (example)", T(-30)],
    ["Nordlys Logistics (example)", "None found (checked)", "Checked partner pages of 12 CBS student organisations and their news page", "", "", "2026", "A.K. (example)", T(-15)],
    ["Fjord Ventures (example)", "None found (checked)", "Checked student organisation partner pages and LinkedIn", "", "", "2026", "A.K. (example)", T(-10)],
    ["Kobber Hotels (example)", "Sponsors similar events", "Hosts local business network events", "https://example.org/kobber-events", "", "2026", "M.S. (example)", T(-12)],
    ["Gamma Software (example)", "None found (checked)", "Checked partner pages and their website", "", "", "2026", "M.S. (example)", T(-50)],
]
for i, v in enumerate(ce_rows, start=3):
    put(ws_ce, i, v)

contacts = [
    ["ProServ (example)", "Head of Employer Branding (example)", "Head of Employer Branding", "Employer branding", "Champion", "Yes", "", "", "Last year's partnership", T(-20), "No", "No"],
    ["MegaBrand A/S (example)", "University Relations Manager (example)", "University Relations Manager", "University relations", "Named", "Unknown", "", "", "LinkedIn", T(-30), "No", "No"],
    ["Fjord Ventures (example)", "Founder (example)", "Founder & CEO", "Founder / CEO (small company)", "Champion", "No", "", "", "Met at our event", T(-10), "No", "No"],
    ["Kobber Hotels (example)", "Marketing Manager (example)", "Marketing Manager", "Marketing", "Named", "Unknown", "", "", "Introduced at a fair", T(-200), "No", "No"],
]
for i, v in enumerate(contacts, start=3):
    put(ws_co, i, v)

opps = [
    ["ProServ (example)", "Recurring sponsor", "Won", "J.H. (example)", "Head of Employer Branding (example)", "No", "Cash", "Strategic package (example)", 40000, 40000, "", "", "", T(-300), T(60), T(30), "Satisfied", "Said the event series met their targets", "", "", "", T(-300), "Paid partner last year"],
    ["MegaBrand A/S (example)", "Recruitment", "Contacted", "M.S. (example)", "University Relations Manager (example)", "No", "Cash", "", "", "", "", "Second follow-up email", T(-3), "", "", "", "Unknown", "", "", "", "", "", "Hires CBS graduates and sponsors student organisations"],
    ["Fjord Ventures (example)", "Speaker", "Meeting", "A.K. (example)", "Founder (example)", "No", "Non-cash", "", "", "", "1 keynote", "Meeting: confirm keynote date", T(4), "", "", "", "Unknown", "", "", "", "", "", "Founder is a strong speaker on a current theme"],
    ["Kobber Hotels (example)", "In-kind", "Proposal sent", "L.M. (example)", "Marketing Manager (example)", "No", "In-kind", "", "", "", "Venue and catering for one evening", "Check venue availability", T(6), "", "", "", "Unknown", "", "", "", "", "", "We need a venue this semester"],
    ["Strandvej Studio (example)", "Event sponsor", "Lost", "M.S. (example)", "", "No", "Cash", "Entry package (example)", 8000, "", "", "", "", "", "", "", "Unknown", "", "No budget", "", "", T(-25), ""],
    ["Søby Foods (example)", "Event sponsor", "Won", "J.H. (example)", "", "No", "Cash", "Entry package (example)", 8000, 8000, "", "", "", T(-30), T(-29), T(300), "Unknown", "", "", "", "", T(-40), ""],
]
for i, v in enumerate(opps, start=3):
    put(ws_op, i, v)

inter = [
    [T(-20), "ProServ (example)", "Head of Employer Branding (example)", "J.H. (example)", "Meeting", "Two-way", "Review of last year's series", "Positive"],
    [T(-17), "MegaBrand A/S (example)", "University Relations Manager (example)", "M.S. (example)", "Email", "Sent by us", "First email about a recruitment event", "No reply"],
    [T(-12), "MegaBrand A/S (example)", "University Relations Manager (example)", "M.S. (example)", "Email", "Sent by us", "Follow-up email", "No reply"],
    [T(-6), "Fjord Ventures (example)", "Founder (example)", "A.K. (example)", "Call", "Two-way", "Agreed to meet about a keynote", "Positive"],
    [T(-9), "Kobber Hotels (example)", "Marketing Manager (example)", "L.M. (example)", "Email", "Received from them", "Interested in hosting; sent proposal", "Positive"],
]
for i, v in enumerate(inter, start=3):
    put(ws_in, i, v)

evd = [
    ["Nordlys Logistics (example)", "F3 Reason to work with us", "About 20 graduate postings for business profiles this year", "Verified fact", "", "https://example.org/nordlys-jobs", "No", "A.K. (example)", T(-15)],
    ["Nordlys Logistics (example)", "F5 Long-term potential", "Recurring graduate hiring each year", "Source-backed inference", "An annual graduate programme implies a recurring hiring need", "https://example.org/nordlys-graduates", "No", "A.K. (example)", T(-15)],
    ["Nordlys Logistics (example)", "General", "Who decides on student partnerships?", "Unknown", "", "", "No", "", ""],
    ["MegaBrand A/S (example)", "F2 Content fit", "CEO spoke about AI and work at a conference", "Verified fact", "", "https://example.org/talk", "Yes", "", ""],
]
for i, v in enumerate(evd, start=3):
    put(ws_evd, i, v)

for ws in wb.worksheets:
    for row_cells in ws.iter_rows():
        for c in row_cells:
            if c.font is None or c.font.name != FONT:
                f = c.font
                c.font = Font(name=FONT, size=f.size or 10, bold=f.bold, italic=f.italic, color=f.color)

wb.save(OUT)
print("saved", OUT)
