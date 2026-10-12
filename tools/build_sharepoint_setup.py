"""Build the SharePoint list setup workbook for the Power Apps version (docs/powerapps-data-model.md).

Run:  python3 tools/build_sharepoint_setup.py  -> powerapps/sharepoint-lists-setup.xlsx

One sheet and one Excel table per SharePoint list. In SharePoint: New -> List -> From Excel,
pick the table, check the column types, create. Column names have no spaces, so SharePoint's
internal names match what the Power Apps formulas use. Every list keeps SharePoint's built-in
Title column as its key (company name, contact name, setting name, ...).

Each table has one fictional example row so SharePoint guesses the column types (text, number,
date) correctly. Delete the example rows in SharePoint after creating the lists.
"""
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = "powerapps/sharepoint-lists-setup.xlsx"
FONT = "Arial"
D = date(2026, 10, 12)

# list name -> [(column, example value)]. "Title" is SharePoint's built-in column.
LISTS = {
    "Companies": [
        ("Title", "Nordlys Logistics (example)"), ("Website", "Nordlys website (example)"), ("CVR", "CVR (example)"),
        ("Industry", "Shipping, logistics & transport"), ("SizeBand", "Unknown"), ("Owner", "A.K. (example)"),
        ("F1", 1), ("F1Reason", "Relevant to supply-chain students"), ("F1Source", ""),
        ("F2", 1), ("F2Reason", "Supply-chain topic, no named speaker yet"), ("F2Source", ""),
        ("F3", 2), ("F3Reason", "About 20 graduate postings this year"), ("F3Source", "Company jobs page (example)"),
        ("F5", 2), ("F5Reason", "Annual graduate programme"), ("F5Source", "Company careers page (example)"),
        ("FitReviewedOn", D), ("IntroPath", "CBS alumnus in their talent team, via A.K. (agreed)"),
        ("SpecificNeed", ""), ("PathInProgress", "No"), ("PathAction", ""), ("PathStartedOn", D),
        ("Notes", ""), ("DoNotContact", "No"), ("DoNotContactReason", ""),
    ],
    "CommercialEvidence": [
        ("Title", "Nordlys Logistics (example)"), ("EvidenceType", "None found (checked)"),
        ("Description", "Checked partner pages of 12 CBS student organisations"), ("SourceUrl", ""),
        ("AmountDKK", 0), ("ObservedYear", "2026"), ("RecordedBy", "A.K. (example)"), ("RecordedOn", D),
    ],
    "Contacts": [
        ("Title", "Mette R. (example)"), ("Company", "Nordlys Logistics (example)"),
        ("JobTitle", "Talent Acquisition Manager"), ("RoleCategory", "Talent acquisition"),
        ("Relationship", "Named"), ("CbsAlumnus", "Unknown"), ("WorkEmail", ""), ("LinkedInUrl", ""),
        ("Source", "Company team page"), ("VerifiedOn", D), ("LeftCompany", "No"), ("DoNotContact", "No"),
    ],
    "Opportunities": [
        ("Title", "Nordlys Logistics (example)"), ("OppType", "Recruitment"), ("Stage", "Planned"),
        ("Owner", "A.K. (example)"), ("Contact", "Mette R. (example)"), ("IsRenewal", "No"),
        ("ValueKind", "Non-cash"), ("Package", ""), ("ProposedDKK", 0), ("AgreedDKK", 0),
        ("NonCashDescription", "Case workshop"), ("NextAction", "Ask alumnus for an introduction"),
        ("NextActionDue", D), ("StartDate", D), ("EndDate", D), ("RenewalDate", D),
        ("Satisfaction", "Unknown"), ("SatisfactionReason", ""), ("LostReason", ""),
        ("NurtureReason", ""), ("NurtureUntil", D), ("ClosedOn", D),
        ("WhyThisType", "They hire ~20 business graduates a year"),
    ],
    "Interactions": [
        ("Title", "Nordlys Logistics (example)"), ("InteractionDate", D), ("Contact", "Mette R. (example)"),
        ("TeamMember", "A.K. (example)"), ("Channel", "Email"), ("Direction", "Sent by us"),
        ("Summary", "Asked about a case workshop"), ("Outcome", "No reply"),
    ],
    "Evidence": [
        ("Title", "Nordlys Logistics (example)"), ("Criterion", "F3 Reason to work with us"),
        ("Statement", "About 20 graduate postings for business profiles this year"),
        ("Label", "Verified fact"), ("Reasoning", ""), ("SourceUrl", "Company jobs page (example)"),
        ("AiDrafted", "No"), ("CheckedBy", "A.K. (example)"), ("CheckedOn", D),
    ],
    "TeamMembers": [
        ("Title", "A.K. (example)"), ("BoardRole", "Partnerships Manager"), ("Email", ""),
        ("Active", "Yes"), ("TermStart", D), ("TermEnd", D),
    ],
    "Packages": [
        ("Title", "Standard package (example)"), ("PriceDKK", 20000),
        ("Includes", "One sponsored event: intro talk, logo, partner report"), ("CapacityPerSemester", 4),
        ("Active", "Yes"),
    ],
    "Themes": [("Title", "Work after AI (example)"), ("ThemeYear", "2026/27"), ("Active", "Yes")],
    "Events": [
        ("Title", "The future of graduate work (example)"), ("EventDate", D), ("Theme", "Work after AI (example)"),
        ("Attendance", 180), ("StudyLineMix", "BSc IB 25%, MSc FIN 15%, other 60%"),
        ("PartnerCompany", ""), ("ReportSentOn", D),
    ],
}

# Settings are key/value rows (all values stored as text; the app converts them).
SETTINGS = [
    ("SemesterName", "Autumn 2026"), ("SemesterStart", "2026-08-01"), ("SemesterEnd", "2027-01-31"),
    ("SemesterCashTargetDKK", "0"), ("FitHigh", "6"), ("AccessHigh", "2"), ("PathCap", "5"),
    ("RenewalDays", "90"), ("NoReplyDays", "10"), ("WarmDays", "365"), ("RecheckDays", "183"),
]

wb = Workbook()
wb.remove(wb.active)
head_font = Font(name=FONT, bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor="1B2A53")


def sheet(name, headers, rows):
    ws = wb.create_sheet(name)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = head_font
        c.fill = head_fill
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(40, len(h) + 6))
    rows = [[("(example)" if v == "" else v) for v in row] for row in rows]
    for r, row in enumerate(rows, start=2):
        for i, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=i, value=None if v == "" else v)
            c.font = Font(name=FONT)
            if isinstance(v, date):
                c.number_format = "yyyy-mm-dd"
    t = Table(displayName=name, ref=f"A1:{get_column_letter(len(headers))}{1 + len(rows)}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)
    ws.freeze_panes = "A2"


readme = wb.create_sheet("Read me")
lines = [
    "SharePoint list setup for the CBS Talks Partnership Desk (Power Apps).",
    "",
    "1. Upload this file to the CBS Talks Team / SharePoint site (Documents).",
    "2. In the site: New -> List -> From Excel -> pick this file -> pick one table (e.g. 'Companies').",
    "3. Before clicking Create, check each column's type in the panel on the right:",
    "   - Number: F1, F2, F3, F5, AmountDKK, ProposedDKK, AgreedDKK, PriceDKK, CapacityPerSemester, Attendance",
    "   - Date and time: columns ending in On, Date, Due, Until, Start or End (e.g. FitReviewedOn, NextActionDue, TermStart)",
    "   - Everything else: Single line of text (change 'Hyperlink', 'Person', 'Choice' or 'Yes/No' to Single line of text)",
    "   If SharePoint says a row was omitted, that's only the example row: check the column types and carry on.",
    "4. Name the list exactly like the table (Companies, CommercialEvidence, Contacts, Opportunities,",
    "   Interactions, Evidence, TeamMembers, Packages, Themes, Events, Settings). Repeat for every table.",
    "5. Delete the example rows (marked '(example)') in each list. Keep all Settings rows and edit their values.",
    "",
    "Column names have no spaces on purpose: the Power Apps formulas use these exact names.",
    "Allowed values (Stage, Outcome, EvidenceType, ...) are enforced by the app's dropdowns,",
    "so SharePoint columns stay plain text. See docs/powerapps-data-model.md in the repository.",
    "All example data is fictional.",
]
for i, t in enumerate(lines, start=1):
    c = readme.cell(row=i, column=1, value=t)
    c.font = Font(name=FONT, bold=(i == 1), size=12 if i == 1 else 10)
    c.alignment = Alignment(wrap_text=False)
readme.column_dimensions["A"].width = 120

for name, cols in LISTS.items():
    sheet(name, [c for c, _ in cols], [[v for _, v in cols]])
sheet("Settings", ["Title", "SettingValue"], [[k, v] for k, v in SETTINGS])

wb.save(OUT)
print("saved", OUT)


# ---------------------------------------------------------------------------------------------
# Test data for the step 3 checkpoint (powerapps/step3-app-skeleton.md): paste into each list's
# grid view. Column order matches the lists. Expected results are in the step 3 guide.
TEST_OUT = "powerapps/step3-test-data.xlsx"
T = date(2026, 10, 3)
COLS = {name: [c for c, _ in cols] for name, cols in LISTS.items()}


def row(name, **vals):
    unknown = set(vals) - set(COLS[name])
    assert not unknown, (name, unknown)
    return [vals.get(c, "") for c in COLS[name]]


TEST = {
    "Companies": [
        row("Companies", Title="ProServ (test)", Owner="A.K. (test)", F1=2, F2=2, F3=2, F5=2, PathInProgress="No", DoNotContact="No"),
        row("Companies", Title="MegaBrand (test)", Owner="A.K. (test)", F1=2, F2=1, F3=1, F5=1, PathInProgress="Yes",
            PathAction="Ask alumni network for an introduction", PathStartedOn=T, DoNotContact="No"),
        row("Companies", Title="Nordlys (test)", Owner="A.K. (test)", F1=1, F2=1, F3=2, F5=2,
            IntroPath="Alumnus in their talent team (test)", PathInProgress="No", DoNotContact="No"),
        row("Companies", Title="Fjord (test)", Owner="A.K. (test)", F1=1, F2=2, F3=1, F5=1, PathInProgress="No", DoNotContact="No"),
        row("Companies", Title="Kobber (test)", Owner="A.K. (test)", F1=0, F2=0, F3=1, F5=1, SpecificNeed="Venue for an autumn event",
            PathInProgress="No", DoNotContact="No"),
        row("Companies", Title="Gamma (test)", Owner="A.K. (test)", F1=0, F2=0, F3=0, F5=0, PathInProgress="No", DoNotContact="No"),
        row("Companies", Title="Vesterled (test)", Owner="A.K. (test)", F1=2, PathInProgress="No", DoNotContact="No"),
    ],
    "CommercialEvidence": [
        row("CommercialEvidence", Title="ProServ (test)", EvidenceType="Previous CBS Talks spend", Description="Paid last year", AmountDKK=40000, ObservedYear="2025"),
        row("CommercialEvidence", Title="MegaBrand (test)", EvidenceType="Sponsors student organisations", Description="Partner of two CBS organisations"),
        row("CommercialEvidence", Title="Nordlys (test)", EvidenceType="None found (checked)", Description="Checked partner pages"),
        row("CommercialEvidence", Title="Fjord (test)", EvidenceType="None found (checked)", Description="Checked partner pages"),
        row("CommercialEvidence", Title="Kobber (test)", EvidenceType="Sponsors similar events", Description="Hosts local business events"),
        row("CommercialEvidence", Title="Gamma (test)", EvidenceType="None found (checked)", Description="Checked partner pages"),
    ],
    "Contacts": [
        row("Contacts", Title="Sofie L. (test)", Company="ProServ (test)", RoleCategory="Employer branding", Relationship="Champion", Source="Test", LeftCompany="No", DoNotContact="No"),
        row("Contacts", Title="Peter N. (test)", Company="MegaBrand (test)", RoleCategory="University relations", Relationship="Named", Source="Test", LeftCompany="No", DoNotContact="No"),
        row("Contacts", Title="Emil B. (test)", Company="Fjord (test)", RoleCategory="Founder / CEO (small company)", Relationship="Champion", Source="Test", LeftCompany="No", DoNotContact="No"),
        row("Contacts", Title="Ida K. (test)", Company="Kobber (test)", RoleCategory="Marketing", Relationship="Named", Source="Test", LeftCompany="No", DoNotContact="No"),
    ],
    "Interactions": [
        row("Interactions", Title="Kobber (test)", InteractionDate=T, Contact="Ida K. (test)", TeamMember="A.K. (test)", Channel="Email",
            Direction="Received from them", Summary="Interested in hosting", Outcome="Positive"),
    ],
}

tw = Workbook()
tw.remove(tw.active)
info = tw.create_sheet("Read me")
for i, t in enumerate([
    "Step 3 test data (fictional). For each tab: open the list with the same name in SharePoint,",
    "click 'Edit in grid view', copy the rows below the header (not the header) and paste them into the first empty row.",
    "Expected results are in powerapps/step3-app-skeleton.md. Delete all '(test)' rows when the checkpoint passes.",
], start=1):
    info.cell(row=i, column=1, value=t).font = Font(name=FONT)
info.column_dimensions["A"].width = 120
for name, rows in TEST.items():
    ws = tw.create_sheet(name)
    for i, h in enumerate(COLS[name], start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = head_font
        c.fill = head_fill
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(30, len(h) + 4))
    for r, vals in enumerate(rows, start=2):
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=None if v == "" else v)
            c.font = Font(name=FONT)
            if isinstance(v, date):
                c.number_format = "yyyy-mm-dd"
tw.save(TEST_OUT)
print("saved", TEST_OUT)
